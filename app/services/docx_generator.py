from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.services.grade_service import GradeResult

logger = logging.getLogger(__name__)


@dataclass
class StatementStudent:
    number: int
    full_name: str
    percent: Decimal
    grade: GradeResult
    individual_number: str = ""
    test_started: str = ""
    note: str = ""
    source_row: int | None = None


@dataclass(frozen=True)
class StatementData:
    subject: str
    group: str
    specialty: str
    teacher: str
    students: list[StatementStudent]
    include_percent: bool = False
    mode: str = "ИФО"
    semester: str = ""
    academic_year: str = ""
    control_form: str = ""
    exam_date: str = ""
    department_head: str = ""


class DocxGenerationError(Exception):
    pass


class DocxGenerator:
    def __init__(self, template_path: str | Path, keep_empty_rows: bool = True):
        self.template_path = Path(template_path)
        self.keep_empty_rows = keep_empty_rows

    def generate(self, data: StatementData, output_path: str | Path) -> Path:
        if not self.template_path.exists():
            raise DocxGenerationError(f"Не найден шаблон Word: {self.template_path}")

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)

        try:
            document = Document(self.template_path)
            self._replace_fields(document, data)
            self._fill_students_table(document, data.students, data.include_percent)
            document.save(output)
            logger.info("Сформирована ведомость: %s", output)
            return output
        except Exception as exc:
            raise DocxGenerationError(f"Не удалось сформировать Word-документ: {exc}") from exc

    def _replace_fields(self, document: Document, data: StatementData) -> None:
        stats = self._grade_stats(data.students)
        replacements = {
            "По предмету: уточняется программой": f"По предмету: {data.subject}",
            "Группа: уточняется программой": f"Группа: {data.group}",
            "Специальность: уточняется программой": f"Специальность: {data.specialty}",
            "Преподаватель: уточняется программой": f"Преподаватель: {data.teacher}",
            "{{subject}}": data.subject,
            "{{group}}": data.group,
            "{{specialty}}": data.specialty,
            "{{teacher}}": data.teacher,
            "{{semester}}": data.semester,
            "{{academic_year}}": data.academic_year,
            "{{control_form}}": data.control_form,
            "{{exam_date}}": data.exam_date,
            "{{department_head}}": data.department_head,
            "{{students_count}}": str(len(data.students)),
            "{{grade_5_count}}": str(stats.get(5, 0)),
            "{{grade_4_count}}": str(stats.get(4, 0)),
            "{{grade_3_count}}": str(stats.get(3, 0)),
            "{{grade_2_count}}": str(stats.get(2, 0)),
            "{{absent_count}}": "0",
        }

        for paragraph in document.paragraphs:
            self._replace_in_paragraph(paragraph, replacements)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, replacements)

    @staticmethod
    def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
        original = paragraph.text
        new_text = original
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text == original:
            return

        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)

    def _fill_students_table(
        self,
        document: Document,
        students: list[StatementStudent],
        include_percent: bool,
    ) -> None:
        table = self._find_students_table(document)
        if table is None:
            raise DocxGenerationError("В шаблоне не найдена таблица студентов")

        header_index = self._find_header_row_index(table)
        if header_index is None:
            raise DocxGenerationError("В таблице студентов не найдена строка заголовков")

        data_row_start = self._find_data_row_start(table, header_index)
        original_data_rows_count = max(0, len(table.rows) - data_row_start)
        required_rows = len(students)
        target_rows = max(required_rows, original_data_rows_count if self.keep_empty_rows else required_rows)

        while len(table.rows) > data_row_start:
            self._remove_row(table, table.rows[-1])
        while len(table.rows) < data_row_start + target_rows:
            table.add_row()

        for row_idx in range(data_row_start, data_row_start + target_rows):
            row = table.rows[row_idx]
            self._clear_row(row)
            student_index = row_idx - data_row_start
            if student_index < len(students):
                student = students[student_index]
                self._write_cell(row.cells[0], str(student.number), align_center=True)
                self._write_cell(row.cells[1], student.full_name)
                if self._is_split_grade_table(row):
                    self._write_cell(row.cells[2], str(student.grade.grade_number), align_center=True)
                    self._write_cell(row.cells[3], student.grade.grade_full, align_center=True)
                    self._write_cell(row.cells[4], "")
                else:
                    self._write_cell(row.cells[2], self._format_grade(student, include_percent), align_center=True)
                    self._write_cell(row.cells[3], "")

    @staticmethod
    def _format_grade(student: StatementStudent, include_percent: bool) -> str:
        if include_percent:
            return student.grade.formatted_grade
        return f"{student.grade.grade_number} ({student.grade.grade_short})"

    @staticmethod
    def _grade_stats(students: list[StatementStudent]) -> dict[int, int]:
        stats: dict[int, int] = {}
        for student in students:
            stats[student.grade.grade_number] = stats.get(student.grade.grade_number, 0) + 1
        return stats

    @staticmethod
    def _is_split_grade_table(row) -> bool:
        return len(row.cells) >= 5

    @staticmethod
    def _find_students_table(document: Document):
        for table in document.tables:
            text = " ".join(cell.text.lower() for row in table.rows for cell in row.cells)
            if "фамилия" in text and "оценка" in text and "подпись" in text:
                return table
        return document.tables[0] if document.tables else None

    @staticmethod
    def _find_header_row_index(table) -> int | None:
        for idx, row in enumerate(table.rows):
            row_text = " ".join(cell.text.lower() for cell in row.cells)
            if "фамилия" in row_text and "оценка" in row_text:
                return idx
        return 0 if table.rows else None

    @staticmethod
    def _find_data_row_start(table, header_index: int) -> int:
        next_index = header_index + 1
        if next_index >= len(table.rows):
            return next_index

        next_row_text = " ".join(cell.text.lower() for cell in table.rows[next_index].cells)
        if "цифрой" in next_row_text or "прописью" in next_row_text:
            return next_index + 1
        return next_index

    @staticmethod
    def _remove_row(table, row) -> None:
        tbl = table._tbl
        tr = row._tr
        tbl.remove(tr)

    @staticmethod
    def _clear_row(row) -> None:
        for cell in row.cells:
            cell.text = ""

    @staticmethod
    def _write_cell(cell, text: str, align_center: bool = False) -> None:
        cell.text = text
        for paragraph in cell.paragraphs:
            if align_center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
