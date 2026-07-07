from decimal import Decimal

from docx import Document

from app.services.docx_generator import DocxGenerator, StatementData, StatementStudent
from app.services.grade_service import GradeResult


def test_ofo_template_replaces_fields_and_splits_grade(tmp_path):
    template_path = "app/templates/statement_template_ofo.docx"
    output_path = tmp_path / "ofo_statement.docx"
    student = StatementStudent(
        number=1,
        full_name="Иванов Иван Иванович",
        percent=Decimal("95"),
        grade=GradeResult(Decimal("95"), 5, "отл", "отлично"),
    )
    data = StatementData(
        subject="Безопасность жизнедеятельности",
        group="421",
        specialty="31.02.02 Акушерское дело",
        teacher="Петров Петр Петрович",
        students=[student],
        mode="ОФО",
        semester="7",
        academic_year="2025/2026",
        control_form="диф.зачет",
        exam_date="01.07.2026",
        department_head="Сидорова С.С.",
    )

    DocxGenerator(template_path).generate(data, output_path)

    document = Document(output_path)
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    text = f"{paragraph_text}\n{table_text}"
    assert "Семестр:" in text
    assert "7" in text
    assert "Учебный год:" in text
    assert "2025/2026" in text
    assert "Число студентов на экзамене (зачете)" in text
    assert "Из них:  получивших \"отлично\"" in text

    student_row = document.tables[1].rows[2]
    assert student_row.cells[1].text == "Иванов Иван Иванович"
    assert student_row.cells[2].text == "5"
    assert student_row.cells[3].text == "отлично"
